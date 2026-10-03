from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader


class DocumentTextExtractionError(Exception):
    pass


def extract_txt_text(
    file_path: Path,
) -> str:
    encodings = (
        "utf-8-sig",
        "cp1251",
    )

    for encoding in encodings:
        try:
            return file_path.read_text(
                encoding=encoding
            ).strip()

        except UnicodeDecodeError:
            continue

    raise DocumentTextExtractionError(
        "Не вдалося визначити "
        "кодування TXT-файла."
    )


def extract_pdf_text(
    file_path: Path,
) -> str:
    try:
        reader = PdfReader(
            str(file_path)
        )

        parts: list[str] = []

        for page in reader.pages:
            text = (
                page.extract_text()
                or ""
            ).strip()

            if text:
                parts.append(text)

        return "\n\n".join(
            parts
        ).strip()

    except Exception as exc:
        raise DocumentTextExtractionError(
            "Не вдалося прочитати PDF."
        ) from exc


def extract_docx_text(
    file_path: Path,
) -> str:
    try:
        document = DocxDocument(
            str(file_path)
        )

        parts: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                parts.append(text)

        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()

                    if text:
                        parts.append(text)

        return "\n\n".join(
            parts
        ).strip()

    except Exception as exc:
        raise DocumentTextExtractionError(
            "Не вдалося прочитати DOCX."
        ) from exc


def extract_document_text(
    file_path: Path,
    extension: str,
) -> str:
    extension = extension.lower()

    if extension == ".txt":
        return extract_txt_text(
            file_path
        )

    if extension == ".pdf":
        return extract_pdf_text(
            file_path
        )

    if extension == ".docx":
        return extract_docx_text(
            file_path
        )

    raise DocumentTextExtractionError(
        "Формат документа "
        "не підтримується."
    )